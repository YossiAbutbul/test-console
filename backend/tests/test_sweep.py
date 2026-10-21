"""Sweep plan validation and Excel export."""

from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from backend.sweep.export import (
    BEST_HEADERS, HEADERS, build_workbook, parse_workbook, rebuild_workbook,
)
from backend.sweep.models import SweepBlock
from backend.sweep.models import ResultRow, SweepConfig


def col(name: str) -> int:
    """1-based column index of a header, so tests survive column insertions."""
    return HEADERS.index(name) + 1


def make_row(idx: int, *, measured: float | None, hp: int = 1, duty: int = 1) -> ResultRow:
    return ResultRow(
        idx=idx,
        freq_hz=902_300_000,
        power_dbm_setting=14,
        pa_duty_cycle=duty,
        hp_max=hp,
        tx_power_dbm=measured,
        current_a=0.1234,
        voltage_v=3.6,
        tx_hex="",
        rx_hex="",
        ok=True,
        status=0,
        t_ms=0,
    )


class TestSweepConfig:
    def test_total_steps_is_the_cartesian_product(self) -> None:
        config = SweepConfig(
            freq_hz=902_300_000,
            power_values=[1, 2, 3],
            duty_values=[1, 2],
            hp_values=[1, 2, 3, 4],
        )
        assert config.total_steps == 24

    def test_defaults_cover_the_full_spec_ranges(self) -> None:
        # 22 powers x 4 duty cycles x 7 hp values — all 1-based, no zero.
        assert SweepConfig(freq_hz=902_300_000).total_steps == 22 * 4 * 7 == 616

    def test_path_loss_defaults_to_no_correction(self) -> None:
        assert SweepConfig(freq_hz=902_300_000).path_loss_db == 0.0

    @pytest.mark.parametrize(
        ("kwargs", "match"),
        [
            ({"power_values": [0]}, "Power out of range"),
            ({"power_values": [23]}, "Power out of range"),
            # Zero hangs the DUT rather than being refused by it, so it has to
            # be stopped here — on all three axes.
            ({"duty_values": [0]}, "PaDutyCycle out of range"),
            ({"duty_values": [5]}, "PaDutyCycle out of range"),
            ({"hp_values": [0]}, "HpMax out of range"),
            ({"hp_values": [8]}, "HpMax out of range"),
        ],
    )
    def test_validate_ranges_rejects_values_the_dut_would_refuse(
        self, kwargs: dict, match: str
    ) -> None:
        config = SweepConfig(freq_hz=902_300_000, **kwargs)
        with pytest.raises(ValueError, match=match):
            config.validate_ranges()

    def test_validate_ranges_accepts_the_boundaries(self) -> None:
        SweepConfig(
            freq_hz=0, power_values=[1, 22], duty_values=[1, 4], hp_values=[1, 7]
        ).validate_ranges()

    def test_the_defaults_never_contain_a_zero(self) -> None:
        cfg = SweepConfig(freq_hz=902_300_000)
        assert 0 not in cfg.power_values
        assert 0 not in cfg.duty_values
        assert 0 not in cfg.hp_values


class TestBuildWorkbook:
    def test_all_sheet_holds_every_row_in_sweep_order(self) -> None:
        rows = [make_row(i, measured=10.0 + i) for i in range(3)]
        wb = load_workbook(BytesIO(build_workbook(rows)))

        assert wb.sheetnames[0] == "All"
        sheet = wb["All"]
        assert [c.value for c in sheet[1]] == HEADERS
        assert sheet.max_row == len(rows) + 1

    def test_one_sheet_per_measured_power_bucket_highest_first(self) -> None:
        rows = [
            make_row(0, measured=20.4),   # rounds to 20
            make_row(1, measured=19.6),   # rounds to 20
            make_row(2, measured=14.0),
        ]
        wb = load_workbook(BytesIO(build_workbook(rows)))

        assert wb.sheetnames == ["All", "Best settings", "20 dBm", "14 dBm"]
        assert wb["20 dBm"].max_row == 3

    def test_bucket_rows_run_from_closest_to_nominal_outwards(self) -> None:
        rows = [make_row(0, measured=20.4), make_row(1, measured=20.1)]
        wb = load_workbook(BytesIO(build_workbook(rows)))

        measured_column = [wb["20 dBm"].cell(row=r, column=5).value for r in (2, 3)]
        assert measured_column == [20.1, 20.4]

    def test_unmeasured_and_under_range_rows_get_no_bucket(self) -> None:
        rows = [
            make_row(0, measured=None),
            make_row(1, measured=-90.0),   # sensor under-range sentinel
            make_row(2, measured=14.0),
        ]
        wb = load_workbook(BytesIO(build_workbook(rows)))

        assert wb.sheetnames == ["All", "Best settings", "14 dBm"]
        # They are still present in All — the run happened, it just didn't read.
        assert wb["All"].max_row == 4

    def test_current_is_exported_in_milliamps(self) -> None:
        wb = load_workbook(BytesIO(build_workbook([make_row(0, measured=14.0)])))
        assert wb["All"].cell(row=2, column=col("CC [mA]")).value == 123.4

    def test_supply_voltage_is_exported_beside_the_current(self) -> None:
        wb = load_workbook(BytesIO(build_workbook([make_row(0, measured=14.0)])))
        assert wb["All"].cell(row=2, column=col("V [V]")).value == 3.6

    def test_a_row_with_no_voltage_leaves_the_cell_empty(self) -> None:
        r = make_row(0, measured=14.0)
        r.voltage_v = None
        wb = load_workbook(BytesIO(build_workbook([r])))
        assert wb["All"].cell(row=2, column=col("V [V]")).value is None

    def test_hp_and_duty_are_exported_as_integers(self) -> None:
        """They used to be "0x07" text, which Excel cannot sort or filter
        numerically. The importer still reads the hex form (see the old-layout
        test below)."""
        wb = load_workbook(BytesIO(build_workbook([make_row(0, measured=14.0, hp=7, duty=3)])))
        sheet = wb["All"]
        assert sheet.cell(row=2, column=col("HP Max")).value == 7
        assert sheet.cell(row=2, column=col("PA DC")).value == 3
        assert sheet.cell(row=2, column=col("HP Max")).number_format == "0"

    def test_the_number_column_is_the_sweep_index_in_every_sheet(self) -> None:
        """A bucket sheet reorders its rows, so a per-sheet counter could not be
        matched back to the All sheet or to the app's table."""
        rows = [
            make_row(0, measured=14.4),   # bucket 14, 0.4 away
            make_row(1, measured=3.0),    # bucket 3
            make_row(2, measured=14.1),   # bucket 14, 0.1 away — sorts first
        ]
        wb = load_workbook(BytesIO(build_workbook(rows)))

        assert [r[col("#") - 1].value for r in wb["All"].iter_rows(min_row=2)] == [1, 2, 3]
        # Bucket 14 holds rows 0 and 2, closest-to-nominal first — but they keep
        # their sweep numbers rather than being renumbered 1, 2.
        assert [r[col("#") - 1].value for r in wb["14 dBm"].iter_rows(min_row=2)] == [3, 1]

    def test_raw_reading_is_exported_so_a_wrong_path_loss_can_be_undone(self) -> None:
        r = make_row(0, measured=34.5)
        r.tx_power_dbm_raw = 14.0
        wb = load_workbook(BytesIO(build_workbook([r])))
        sheet = wb["All"]
        assert sheet.cell(row=2, column=col("Measured [dBm]")).value == 34.5
        assert sheet.cell(row=2, column=col("Raw [dBm]")).value == 14.0

    def test_a_failed_row_says_why_in_the_file(self) -> None:
        r = make_row(0, measured=None)
        r.current_a = None
        r.ok = False
        r.error = "DUT rejected the command (status=255)"
        wb = load_workbook(BytesIO(build_workbook([r])))
        assert "255" in str(wb["All"].cell(row=2, column=col("Status")).value)

    def test_run_sheet_records_the_config_right_after_all(self) -> None:
        cfg = SweepConfig(
            freq_hz=902_300_000,
            power_values=[1, 2, 3],
            duty_values=[1],
            hp_values=[5],
            settle_ms=450,
            path_loss_db=20.5,
        )
        wb = load_workbook(
            BytesIO(build_workbook([make_row(0, measured=14.0)], cfg, 1_700_000_000.0, None))
        )

        assert wb.sheetnames[:2] == ["All", "Run"], "Run must be findable, not behind the buckets"
        pairs = {r[0].value: r[1].value for r in wb["Run"].iter_rows(min_row=2)}
        assert pairs["Path loss [dB]"] == 20.5
        assert pairs["Settle [ms]"] == 450
        assert pairs["Frequency [MHz]"] == 902.3
        assert pairs["Power Set values"] == "1–3 (3 values)"
        assert pairs["Finished"] == "—", "an unfinished run must not invent an end time"
        assert pairs["Current measured"] == "Yes"

    def test_best_settings_sheet_sits_after_run_and_before_the_buckets(self) -> None:
        cfg = SweepConfig(freq_hz=902_300_000, power_values=[14], duty_values=[1], hp_values=[1])
        wb = load_workbook(BytesIO(build_workbook([make_row(0, measured=14.0)], cfg)))
        assert wb.sheetnames == ["All", "Run", "Best settings", "14 dBm"]

    def test_best_settings_picks_the_cheapest_row_per_level_strongest_first(self) -> None:
        def at(idx: int, measured: float, amps: float | None, hp: int) -> ResultRow:
            r = make_row(idx, measured=measured, hp=hp)
            r.current_a = amps
            return r

        rows = [
            at(0, 14.2, 0.040, hp=1),
            at(1, 13.8, 0.021, hp=2),   # same level, cheaper
            at(2, 14.0, None, hp=3),    # no reading: never the recommendation
            at(3, 20.0, 0.090, hp=4),
            at(4, 8.0, None, hp=5),     # level reached, nothing to recommend
        ]
        ws = load_workbook(BytesIO(build_workbook(rows)))["Best settings"]
        assert [c.value for c in ws[1]] == BEST_HEADERS

        def best(name: str) -> int:
            return BEST_HEADERS.index(name) + 1

        table = [
            (ws.cell(row=r, column=1).value,
             ws.cell(row=r, column=best("#")).value,
             ws.cell(row=r, column=best("CC [mA]")).value,
             ws.cell(row=r, column=best("Status")).value,
             ws.cell(row=r, column=best("Rows at level")).value)
            for r in range(2, ws.max_row + 1)
        ]
        assert table == [
            (20, 4, 90.0, "ok", 1),
            (14, 2, 21.0, "ok", 3),
            (8, None, None, "no current reading", 1),
        ]

    def test_run_sheet_is_omitted_when_no_config_is_supplied(self) -> None:
        wb = load_workbook(BytesIO(build_workbook([make_row(0, measured=14.0)])))
        assert "Run" not in wb.sheetnames

    def test_empty_result_set_still_produces_a_valid_workbook(self) -> None:
        wb = load_workbook(BytesIO(build_workbook([])))
        assert wb.sheetnames == ["All"]
        assert [c.value for c in wb["All"][1]] == HEADERS


class TestRebuildWorkbook:
    """Re-exporting an imported file: same rows, current sheets, its own Run."""

    @staticmethod
    def _old_export(cfg: SweepConfig | None) -> bytes:
        """A workbook as written before the Best settings sheet existed."""
        wb = load_workbook(BytesIO(build_workbook(
            [make_row(0, measured=20.2), make_row(1, measured=14.0)], cfg, 1_700_000_000.0, None,
        )))
        del wb["Best settings"]
        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def test_an_old_export_gains_the_best_settings_sheet(self) -> None:
        cfg = SweepConfig(freq_hz=902_300_000, power_values=[14], duty_values=[1], hp_values=[1])
        wb = load_workbook(BytesIO(rebuild_workbook(self._old_export(cfg))))
        assert wb.sheetnames == ["All", "Run", "Best settings", "20 dBm", "14 dBm"]
        assert wb["Best settings"].max_row == 3

    def test_the_run_sheet_is_carried_over_unchanged(self) -> None:
        cfg = SweepConfig(
            freq_hz=902_300_000, power_values=[14], duty_values=[1], hp_values=[1],
            settle_ms=450, path_loss_db=20.5,
        )
        source = self._old_export(cfg)

        def pairs(data: bytes) -> list[tuple]:
            return [tuple(r) for r in load_workbook(BytesIO(data))["Run"].iter_rows(values_only=True)]

        # Copied, not rebuilt: a re-export must not invent a different
        # path loss, row count or end time for a run it did not make.
        assert pairs(rebuild_workbook(source)) == pairs(source)

    def test_rows_survive_the_round_trip(self) -> None:
        source = self._old_export(None)
        assert parse_workbook(rebuild_workbook(source)) == parse_workbook(source)

    def test_a_file_without_a_run_sheet_gets_none(self) -> None:
        wb = load_workbook(BytesIO(rebuild_workbook(self._old_export(None))))
        assert "Run" not in wb.sheetnames

    def test_refuses_a_workbook_that_is_not_ours(self) -> None:
        wb = Workbook()
        buf = BytesIO()
        wb.save(buf)
        with pytest.raises(ValueError, match="not a sweep workbook"):
            rebuild_workbook(buf.getvalue())


class TestParseWorkbook:
    """`parse_workbook` is the inverse of `build_workbook`; they share HEADERS,
    so these tests are what stop the two drifting apart."""

    def test_round_trips_every_field_the_workbook_carries(self) -> None:
        cfg = SweepConfig(freq_hz=902_300_000, settle_ms=450, path_loss_db=20.5)
        original = make_row(0, measured=18.89, hp=5, duty=1)
        original.tx_power_dbm_raw = -1.61
        back = parse_workbook(build_workbook([original], cfg))

        assert len(back) == 1
        r = back[0]
        assert r.idx == 0
        assert (r.hp_max, r.pa_duty_cycle) == (5, 1)
        assert r.power_dbm_setting == original.power_dbm_setting
        assert r.tx_power_dbm == pytest.approx(18.89)
        assert r.tx_power_dbm_raw == pytest.approx(-1.61)
        # Written as mA, restored as amps.
        assert r.current_a == pytest.approx(original.current_a, abs=1e-6)
        assert r.voltage_v == pytest.approx(3.6)
        assert r.ok is True

    def test_recovers_the_frequency_from_the_run_sheet(self) -> None:
        cfg = SweepConfig(freq_hz=908_700_000)
        back = parse_workbook(build_workbook([make_row(0, measured=14.0)], cfg))
        assert back[0].freq_hz == 908_700_000

    def test_a_file_without_a_run_sheet_still_imports(self) -> None:
        back = parse_workbook(build_workbook([make_row(0, measured=14.0)]))
        assert len(back) == 1 and back[0].freq_hz == 0

    def test_a_failed_row_comes_back_as_failed(self) -> None:
        r = make_row(0, measured=None)
        r.current_a = None
        r.ok = False
        r.error = "DUT rejected the command (status=255)"
        back = parse_workbook(build_workbook([r]))
        assert back[0].ok is False
        assert back[0].error is not None and "255" in back[0].error

    def test_rejects_a_workbook_that_is_not_ours(self) -> None:
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.title = "Sheet1"
        wb.active.cell(row=1, column=1, value="something else")
        buf = BytesIO()
        wb.save(buf)
        with pytest.raises(ValueError, match="not a sweep workbook"):
            parse_workbook(buf.getvalue())

    def test_rejects_a_workbook_missing_our_columns(self) -> None:
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.title = "All"
        for c, name in enumerate(["#", "Wrong", "Columns"], start=1):
            wb.active.cell(row=1, column=c, value=name)
        buf = BytesIO()
        wb.save(buf)
        with pytest.raises(ValueError, match="missing columns"):
            parse_workbook(buf.getvalue())

    def test_reads_a_file_written_before_the_columns_were_reordered(self) -> None:
        """The old layout put HP Max and PA DC ahead of Power Set.

        Files exported by earlier versions are still on disk and still get
        imported, so the parser matches on the header text rather than on
        position. This is that promise, written down.
        """
        from openpyxl import Workbook
        # Also the old "Power Set [dBm]" spelling, since renamed "Set [dBm]".
        old_headers = [
            "#", "HP Max", "PA DC", "Power Set [dBm]", "Measured [dBm]",
            "Raw [dBm]", "CC [mA]", "V [V]", "Status",
        ]
        wb = Workbook()
        ws = wb.active
        ws.title = "All"
        for c, name in enumerate(old_headers, start=1):
            ws.cell(row=1, column=c, value=name)
        # idx, hp, duty, set, measured, raw, cc, v, status -- in the old order.
        for c, value in enumerate(
            [1, "0x03", "0x02", 14, 13.5, -22.3, 250.0, 3.7, "ok"], start=1,
        ):
            ws.cell(row=2, column=c, value=value)
        buf = BytesIO()
        wb.save(buf)

        back = parse_workbook(buf.getvalue())
        assert len(back) == 1
        r = back[0]
        # Each value has to land on its own field, not on the one that now
        # occupies its old column -- power and hp_max are the pair that would
        # silently swap if position were still used.
        assert r.hp_max == 3
        assert r.pa_duty_cycle == 2
        assert r.power_dbm_setting == 14
        assert r.tx_power_dbm == 13.5
        assert r.tx_power_dbm_raw == -22.3
        assert r.current_a is not None and abs(r.current_a - 0.25) < 1e-9
        assert r.voltage_v == 3.7
        assert r.ok is True


class TestSweepBlocks:
    """Multi-range plans: several cross-products in one run."""

    def test_no_blocks_means_the_legacy_single_cross_product(self) -> None:
        # Every client that predates blocks sends only the three lists, and
        # must keep behaving exactly as it did.
        cfg = SweepConfig(
            freq_hz=902_300_000,
            power_values=[1, 2, 3], duty_values=[1, 2], hp_values=[1, 2, 3, 4],
        )
        assert cfg.blocks == []
        assert cfg.total_steps == 3 * 2 * 4
        blocks = cfg.effective_blocks
        assert len(blocks) == 1
        assert blocks[0].power_values == [1, 2, 3]
        assert blocks[0].duty_values == [1, 2]
        assert blocks[0].hp_values == [1, 2, 3, 4]

    def test_steps_are_the_sum_of_each_block(self) -> None:
        cfg = SweepConfig(
            freq_hz=902_300_000,
            blocks=[
                SweepBlock(power_values=list(range(1, 13)), duty_values=[2, 3, 4], hp_values=[1]),
                SweepBlock(power_values=list(range(13, 23)), duty_values=[1], hp_values=[1, 2]),
            ],
        )
        assert cfg.total_steps == 12 * 3 * 1 + 10 * 1 * 2

    def test_blocks_win_over_the_flat_lists(self) -> None:
        # Both present is a client that sent blocks without clearing the
        # legacy fields; the explicit plan is the one that was meant.
        cfg = SweepConfig(
            freq_hz=902_300_000,
            power_values=list(range(1, 23)),
            blocks=[SweepBlock(power_values=[5], duty_values=[1], hp_values=[1])],
        )
        assert cfg.total_steps == 1

    def test_every_block_is_range_checked(self) -> None:
        # A bad value in the second block must not slip through because the
        # first one was fine.
        cfg = SweepConfig(
            freq_hz=902_300_000,
            blocks=[
                SweepBlock(power_values=[1], duty_values=[1], hp_values=[1]),
                SweepBlock(power_values=[23], duty_values=[1], hp_values=[1]),
            ],
        )
        with pytest.raises(ValueError, match="Power out of range"):
            cfg.validate_ranges()

    def test_the_run_sheet_lists_each_block(self) -> None:
        from openpyxl import load_workbook
        cfg = SweepConfig(
            freq_hz=902_300_000,
            blocks=[
                SweepBlock(power_values=[1, 2], duty_values=[2], hp_values=[1]),
                SweepBlock(power_values=[5], duty_values=[1], hp_values=[1, 2]),
            ],
        )
        wb = load_workbook(BytesIO(build_workbook([make_row(0, measured=14.0)], cfg)))
        run = {r[0]: r[1] for r in wb["Run"].iter_rows(min_row=2, values_only=True) if r[0]}
        assert "Block 1" in run and "Block 2" in run
        assert str(run["Block 1"]).startswith("Power 1, 2")
        assert "PA DC 2" in str(run["Block 1"])
        assert run["Steps planned"] == 2 * 1 * 1 + 1 * 1 * 2
        # The flat per-axis lines would be a lie here -- there is no single
        # span that describes both blocks.
        assert "Power Set values" not in run

    def test_a_single_block_keeps_the_flat_run_sheet(self) -> None:
        from openpyxl import load_workbook
        cfg = SweepConfig(freq_hz=902_300_000, power_values=[1, 2], duty_values=[1], hp_values=[1])
        wb = load_workbook(BytesIO(build_workbook([make_row(0, measured=14.0)], cfg)))
        run = {r[0]: r[1] for r in wb["Run"].iter_rows(min_row=2, values_only=True) if r[0]}
        assert "Power Set values" in run
        assert "Block 1" not in run


class TestWorkbookFormatting:
    """Display only: the cells keep full precision, the format rounds."""

    def test_readings_are_stored_unrounded_and_shown_to_two_decimals(self) -> None:
        r = make_row(0, measured=14.123456)
        r.tx_power_dbm_raw = -6.376544
        r.current_a = 0.18567891
        ws = load_workbook(BytesIO(build_workbook([r])))["All"]
        meas = ws.cell(row=2, column=col("Measured [dBm]"))
        cc = ws.cell(row=2, column=col("CC [mA]"))
        assert meas.value == 14.123456
        assert cc.value == pytest.approx(185.67891)
        assert meas.number_format == "0.00"
        assert cc.number_format == "0.00"
        assert ws.cell(row=2, column=col("Raw [dBm]")).number_format == "0.00"
        assert ws.cell(row=2, column=col("V [V]")).number_format == "0.00"
        assert ws.cell(row=2, column=col("#")).number_format == "0"
        assert ws.cell(row=2, column=col("Set [dBm]")).number_format == "0"

    def test_data_sheets_are_frozen_filtered_and_styled(self) -> None:
        rows = [make_row(0, measured=14.0), make_row(1, measured=14.2)]
        wb = load_workbook(BytesIO(build_workbook(rows)))
        for name in ("All", "Best settings", "14 dBm"):
            ws = wb[name]
            assert ws.freeze_panes == "B2", name
            assert ws.auto_filter.ref and ws.auto_filter.ref.startswith("A1:"), name
            assert ws.cell(row=1, column=1).font.b, name
            assert ws.cell(row=1, column=1).fill.fgColor.rgb.endswith("DDE5F0"), name

    def test_failed_status_is_flagged_and_current_is_colour_scaled(self) -> None:
        wb = load_workbook(BytesIO(build_workbook(
            [make_row(0, measured=14.0), make_row(1, measured=14.3)],
        )))

        def rules(ws):
            return [(str(rng.sqref), r.type) for rng in ws.conditional_formatting for r in rng.rules]

        status = chr(ord("A") + col("Status") - 1)
        cc_best = chr(ord("A") + BEST_HEADERS.index("CC [mA]"))
        cc = chr(ord("A") + col("CC [mA]") - 1)
        assert (f"{status}2:{status}3", "expression") in rules(wb["All"])
        assert (f"{cc}2:{cc}3", "colorScale") in rules(wb["14 dBm"])
        # One level, so a one-cell range (openpyxl writes it as "X2").
        assert (f"{cc_best}2", "colorScale") in rules(wb["Best settings"])
        # All mixes drive levels, so ranking its currents would mean nothing.
        assert not any(t == "colorScale" for _, t in rules(wb["All"]))

    def test_run_sheet_carries_the_rig_and_a_summary(self) -> None:
        from backend.sweep.models import RigInfo
        cfg = SweepConfig(freq_hz=902_300_000, power_values=[14], duty_values=[1], hp_values=[1])
        bad = make_row(1, measured=None)
        bad.ok = False
        bad.error = "tx status=3"
        rig = RigInfo(
            dut_mac="80:E1:27:20:76:DD", dut_name="CATM2-DUT",
            power_sensor_idn="PWR-8GHS,12345,A1", power_sensor_address="12345",
            dc_analyzer_idn="Keysight N6705B", dc_analyzer_address="USB0::X::INSTR",
            dc_channel=3, operator="RF", notes="after rework",
        )
        wb = load_workbook(BytesIO(build_workbook(
            [make_row(0, measured=14.2), bad], cfg, 1_700_000_000.0, 1_700_000_065.0, rig,
        )))
        run = {r[0]: r[1] for r in wb["Run"].iter_rows(min_row=2, values_only=True) if r[0]}
        assert run["DUT MAC"] == "80:E1:27:20:76:DD"
        assert run["DUT device"] == "CATM2-DUT"
        assert run["Power sensor"] == "PWR-8GHS,12345,A1"
        assert run["DC analyzer address"] == "USB0::X::INSTR"
        assert run["DC channel"] == 3
        assert run["Operator"] == "RF" and run["Notes"] == "after rework"
        assert (run["Points ok"], run["Points failed"]) == (1, 1)
        assert run["Measured min [dBm]"] == run["Measured max [dBm]"] == 14.2
        assert run["Duration"] == "0:01:05"
        assert run["PA mode"] == "Off (0)", "one PA-mode spelling across CSV and xlsx"
        assert run["App build"]
        assert "Run" in wb.sheetnames and run["Path loss [dB]"] == 0.0

    def test_a_run_sheet_without_rig_details_says_so(self) -> None:
        cfg = SweepConfig(freq_hz=902_300_000, power_values=[14], duty_values=[1], hp_values=[1])
        wb = load_workbook(BytesIO(build_workbook([make_row(0, measured=14.0)], cfg)))
        run = {r[0]: r[1] for r in wb["Run"].iter_rows(min_row=2, values_only=True) if r[0]}
        assert run["Power sensor"] == "—"
        assert "Operator" not in run

    def test_an_export_re_imports_with_full_precision(self) -> None:
        cfg = SweepConfig(freq_hz=902_312_500)
        r = make_row(0, measured=14.123456)
        back = parse_workbook(build_workbook([r], cfg))[0]
        assert back.tx_power_dbm == 14.123456
        assert back.freq_hz == 902_312_500
        assert back.current_a == pytest.approx(0.1234)


class TestExportFilename:
    def test_mac_slug_test_slug_and_timestamp(self) -> None:
        import re
        from backend.workbook import export_filename
        name = export_filename("80:e1:27:20:76:dd", "power-sweep", "xlsx")
        assert re.fullmatch(r"80E1272076DD_power-sweep_\d{4}-\d{2}-\d{2}T\d{2}-\d{2}-\d{2}\.xlsx", name)

    def test_no_dut(self) -> None:
        from backend.workbook import export_filename
        assert export_filename(None, "load-pull", "xlsx").startswith("noDUT_load-pull_")
