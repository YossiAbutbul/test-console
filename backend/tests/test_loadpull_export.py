"""Load Pull workbook layout.

The split into a frequency tab per frequency, and a table per power inside it,
is the whole point of this export - a contour only means something at one
frequency and one drive level - so the shape is pinned here.
"""
from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook

from backend.loadpull import (
    HEADERS, TABLE_HEADERS, LoadPullMeta, LoadPullRow, build_workbook, parse_workbook,
)
from backend.loadpull.models import PathLossPoint


def row(pos_mm: float, freq: float | None, power: float | None, **over) -> LoadPullRow:
    base = dict(
        pos_pulses=int(pos_mm * 400),
        pos_mm=pos_mm,
        freq_mhz=freq,
        power_dbm_setting=power,
        power_dbm=12.5,
        current_a=0.0225,
        r_ohm=48.2,
        x_ohm=-3.1,
        s11_db=-14.8,
    )
    base.update(over)
    return LoadPullRow(**base)


def wb_of(rows):
    return load_workbook(BytesIO(build_workbook(rows)))


class TestLoadPullWorkbook:
    def test_all_sheet_holds_every_row_in_order(self) -> None:
        rows = [row(0, 902.3, 0), row(1, 915, 14), row(2, 902.3, 14)]
        wb = wb_of(rows)

        assert wb.sheetnames[0] == "All"
        sheet = wb["All"]
        assert [c.value for c in sheet[1]] == HEADERS
        assert sheet.max_row == len(rows) + 1
        # Untouched order: the sweep visits positions in a deliberate sequence.
        assert [sheet.cell(row=r, column=2).value for r in (2, 3, 4)] == [0, 1, 2]

    def test_one_tab_per_frequency_ascending(self) -> None:
        wb = wb_of([row(0, 915, 0), row(1, 902.3, 0)])
        assert wb.sheetnames == ["All", "902.3 MHz", "915 MHz"]

    def test_a_frequency_tab_holds_one_table_per_power(self) -> None:
        wb = wb_of([
            row(0, 902.3, 14), row(1, 902.3, 14),
            row(2, 902.3, 0), row(3, 902.3, 0), row(4, 902.3, 0),
        ])
        col_a = [c.value for c in wb["902.3 MHz"]["A"]]

        # Ascending power, each table titled and separated by a blank row.
        assert col_a.count("Set power: 0 dBm") == 1
        assert col_a.count("Set power: 14 dBm") == 1
        assert col_a.index("Set power: 0 dBm") < col_a.index("Set power: 14 dBm")
        assert col_a.count("#") == 2, "each table repeats its own header"

    def test_frequency_tables_drop_the_columns_their_title_states(self) -> None:
        wb = wb_of([row(0, 902.3, 14)])
        sheet = wb["902.3 MHz"]
        assert [c.value for c in sheet[2]][: len(TABLE_HEADERS)] == TABLE_HEADERS
        assert "Freq [MHz]" not in TABLE_HEADERS
        assert "Set [dBm]" not in TABLE_HEADERS

    def test_current_is_exported_in_milliamps(self) -> None:
        wb = wb_of([row(0, 902.3, 0, current_a=0.1234)])
        assert wb["All"].cell(row=2, column=HEADERS.index("CC [mA]") + 1).value == 123.4

    def test_a_failed_point_says_why(self) -> None:
        wb = wb_of([row(0, 902.3, 0, power_dbm=None, error="tx status=3")])
        assert wb["All"].cell(row=2, column=HEADERS.index("Status") + 1).value == "tx status=3"

    def test_rows_without_a_frequency_are_kept_not_dropped(self) -> None:
        """Rows recorded before a run could sweep, and CSVs exported then."""
        wb = wb_of([row(0, 902.3, 0), row(1, None, None)])
        assert wb.sheetnames == ["All", "902.3 MHz", "Unspecified"]
        assert wb["All"].max_row == 3

    def test_a_whole_number_frequency_reads_without_a_decimal(self) -> None:
        assert "915 MHz" in wb_of([row(0, 915.0, 0)]).sheetnames

    def test_empty_result_set_still_produces_a_valid_workbook(self) -> None:
        wb = wb_of([])
        assert wb.sheetnames == ["All"]
        assert [c.value for c in wb["All"][1]] == HEADERS


class TestParseWorkbook:
    """`parse_workbook` is the inverse of `build_workbook`; they share HEADERS,
    which is what stops the two drifting apart."""

    def test_round_trips_every_field(self) -> None:
        original = row(1.5, 902.3, 14, current_a=0.0221)
        back = parse_workbook(build_workbook([original]))

        assert len(back) == 1
        r = back[0]
        assert (r.pos_mm, r.pos_pulses) == (original.pos_mm, original.pos_pulses)
        assert (r.freq_mhz, r.power_dbm_setting) == (902.3, 14)
        assert r.power_dbm == pytest.approx(12.5)
        # Written as mA, restored as amps.
        assert r.current_a == pytest.approx(0.0221)
        assert (r.r_ohm, r.x_ohm, r.s11_db) == (48.2, -3.1, -14.8)
        assert r.error is None

    def test_only_the_all_sheet_is_read(self) -> None:
        """The frequency tabs repeat the same points, so reading them too would
        import every row twice."""
        rows = [row(0, 902.3, 0), row(1, 915, 14)]
        assert len(parse_workbook(build_workbook(rows))) == len(rows)

    def test_a_failed_row_comes_back_as_failed(self) -> None:
        back = parse_workbook(build_workbook([row(0, 902.3, 0, power_dbm=None, error="tx status=3")]))
        assert back[0].error == "tx status=3"
        assert back[0].power_dbm is None

    def test_rows_without_a_frequency_survive_the_trip(self) -> None:
        back = parse_workbook(build_workbook([row(0, None, None)]))
        assert back[0].freq_mhz is None and back[0].power_dbm_setting is None

    def test_rejects_a_workbook_that_is_not_ours(self) -> None:
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.title = "Sheet1"
        wb.active.cell(row=1, column=1, value="something else")
        buf = BytesIO()
        wb.save(buf)
        with pytest.raises(ValueError, match="not a Load Pull workbook"):
            parse_workbook(buf.getvalue())

    def test_rejects_a_workbook_whose_columns_moved(self) -> None:
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.title = "All"
        for c, name in enumerate(["#", "Wrong", "Columns"], start=1):
            wb.active.cell(row=1, column=c, value=name)
        buf = BytesIO()
        wb.save(buf)
        with pytest.raises(ValueError, match="unexpected columns"):
            parse_workbook(buf.getvalue())


class TestRunSheet:
    def test_run_sheet_records_when_and_with_what(self) -> None:
        meta = LoadPullMeta(
            freq_spec="902.3,915", power_spec="0,14", settle_ms=400, pa_mode=0,
            delta_x_mm=5, zero_pulses=-49732, end_pulses=5868,
            path_loss_default_db=36,
            path_loss_points=[
                PathLossPoint(freq_mhz=902.3, db=20.5, calibrated=True),
                PathLossPoint(freq_mhz=915, db=36, calibrated=False),
            ],
            dut_mac="80:E1:27:20:76:DD",
        )
        wb = load_workbook(BytesIO(build_workbook([row(0, 902.3, 14)], meta)))

        assert wb.sheetnames[:2] == ["All", "Run"], "Run must be findable, not behind the tabs"
        pairs = {r[0]: r[1] for r in wb["Run"].iter_rows(min_row=2, values_only=True) if r[0]}
        assert pairs["Frequencies"] == "902.3,915"
        assert pairs["Powers"] == "0,14"
        assert pairs["Settle [ms]"] == 400
        assert pairs["Points"] == 1
        assert pairs["DUT"] == "80:E1:27:20:76:DD"
        # A timestamp, not a placeholder.
        assert len(str(pairs["Exported"])) == len("2026-08-11 10:41:13")

    def test_an_uncalibrated_path_loss_says_so(self) -> None:
        meta = LoadPullMeta(path_loss_points=[
            PathLossPoint(freq_mhz=915, db=36, calibrated=False),
        ])
        wb = load_workbook(BytesIO(build_workbook([row(0, 915, 0)], meta)))
        pairs = {r[0]: r[1] for r in wb["Run"].iter_rows(min_row=2, values_only=True) if r[0]}
        assert "uncalibrated" in str(pairs["Path loss @ 915 MHz [dB]"])

    def test_the_run_sheet_is_omitted_without_meta(self) -> None:
        assert "Run" not in load_workbook(BytesIO(build_workbook([row(0, 902.3, 0)]))).sheetnames

    def test_a_frequency_tab_cannot_collide_with_the_run_sheet(self) -> None:
        """Sheet titles must stay unique; "Run" is taken before the tabs are made."""
        wb = load_workbook(BytesIO(build_workbook([row(0, 902.3, 0)], LoadPullMeta())))
        assert len(wb.sheetnames) == len(set(wb.sheetnames))


class TestAttenuation:
    """The attenuator is manual, so a run records a whole trombone cycle per
    setting; the workbook is what lets those cycles be compared."""

    def test_att_rows_sort_ascending_within_a_power_table(self) -> None:
        rows = []
        for att in (2, 0, 1):          # deliberately not in order
            for pos in (0, 1):
                rows.append(row(pos, 902.3, 14, att_db=att))
        wb = wb_of(rows)

        data = [r for r in wb["902.3 MHz"].iter_rows(min_row=3, values_only=True)
                if r[0] is not None]
        att_col = TABLE_HEADERS.index("Att [dB]")
        assert [r[att_col] for r in data] == [0, 0, 1, 1, 2, 2]

    def test_position_order_survives_inside_one_att_setting(self) -> None:
        rows = [row(0, 902.3, 14, att_db=0), row(1, 902.3, 14, att_db=0)]
        wb = wb_of(rows)
        data = [r for r in wb["902.3 MHz"].iter_rows(min_row=3, values_only=True)
                if r[0] is not None]
        pos_col = TABLE_HEADERS.index("Pos [mm]")
        assert [r[pos_col] for r in data] == [0, 1]

    def test_rows_without_att_sort_ahead_of_those_with_it(self) -> None:
        """A run that did not sweep attenuation must not be reordered by it."""
        rows = [row(0, 902.3, 14, att_db=5), row(1, 902.3, 14)]
        wb = wb_of(rows)
        data = [r for r in wb["902.3 MHz"].iter_rows(min_row=3, values_only=True)
                if r[0] is not None]
        att_col = TABLE_HEADERS.index("Att [dB]")
        assert [r[att_col] for r in data] == [None, 5]

    def test_att_survives_the_round_trip(self) -> None:
        back = parse_workbook(build_workbook([row(0, 902.3, 14, att_db=7)]))
        assert back[0].att_db == 7

    def test_att_is_kept_on_the_frequency_tables(self) -> None:
        """Frequency and power are stated by the tab and title; attenuation
        varies within a table, so it has to stay a column."""
        assert "Att [dB]" in TABLE_HEADERS
