"""Presentation shared by the exported workbooks (sweep and load pull).

Both exporters used to carry their own header font, their own `_autosize`, and
their own idea of how many decimals a reading deserved. Kept here instead, so a
sweep file and a load-pull file opened side by side look like one app made them.

Numbers are written at full precision and rounded *by the cell format only*.
Rounding in Python made the file lossy: a re-import, or a reader working out a
delta between two cells, got the display value rather than the measurement.
"""
from __future__ import annotations

import subprocess
import sys
import time
from functools import lru_cache
from pathlib import Path
from typing import Any, Iterable, Sequence

from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill("solid", start_color="DDE5F0", end_color="DDE5F0")
#: Section labels on the Run sheet ("Summary", "Rig"): bold like a header, but
#: a lighter fill so they do not read as a second table.
SECTION_FILL = PatternFill("solid", start_color="EEF2F7", end_color="EEF2F7")
BAD_FILL = PatternFill("solid", start_color="FFC7CE", end_color="FFC7CE")
BAD_FONT = Font(color="9C0006")

#: Two decimals is the operator's explicit requirement for every reading.
DECIMAL_FORMAT = "0.00"
INTEGER_FORMAT = "0"
#: Frequencies keep at least two decimals like everything else, but are allowed
#: up to four: a 902.3125 MHz channel shown as 902.31 would name a different
#: channel, which a power reading rounded the same way does not do.
FREQ_FORMAT = "0.00##"

#: Columns whose content is wider than their header. Everything else gets the
#: header's own width plus room for the auto-filter button.
_WIDTHS = {
    "#": 7,
    "Status": 32,
    "Setting": 26,
    "Value": 44,
}
_WIDTH_MIN = 9
#: The auto-filter drop-down sits over the right edge of the header cell, so a
#: width that exactly fits the text hides its last characters behind it.
_FILTER_PAD = 5


def number_format(header: str | None, value: Any) -> str | None:
    """The display format for one cell, chosen by what it holds.

    By value type rather than by a per-column list: a column is either an int
    field or a float field on the row model, so the type already says which it
    is, and a column added later is formatted without anyone remembering to.
    """
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    if isinstance(value, int):
        return INTEGER_FORMAT
    if header is not None and header.startswith("Freq"):
        return FREQ_FORMAT
    return DECIMAL_FORMAT


def write_cell(ws: Worksheet, row: int, col: int, value: Any, header: str | None = None) -> None:
    cell = ws.cell(row=row, column=col, value=value)
    fmt = number_format(header, value)
    if fmt is not None:
        cell.number_format = fmt


def write_header(ws: Worksheet, headers: Sequence[str], row: int = 1, start_col: int = 1) -> None:
    for c, name in enumerate(headers, start=start_col):
        cell = ws.cell(row=row, column=c, value=name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL


def set_widths(ws: Worksheet, headers: Sequence[str], start_col: int = 1) -> None:
    """Fixed widths per header.

    Replaces sizing from the longest value: that measured `str(float)`, so a
    full-precision reading (`-1.6099999999999999`) blew a column out to the
    cap even though the cell only ever displays `-1.61`.
    """
    for c, name in enumerate(headers, start=start_col):
        width = _WIDTHS.get(name, max(len(name) + _FILTER_PAD, _WIDTH_MIN))
        ws.column_dimensions[get_column_letter(c)].width = width


def finish_table(
    ws: Worksheet,
    headers: Sequence[str],
    row_count: int,
    *,
    freeze: str | None = "B2",
    autofilter: bool = True,
) -> None:
    """Freeze, filter and size a flat table whose header is row 1.

    Frozen at B2 by default: the header row and the "#" column, which is what
    lets a row be traced back to the All sheet while scrolling a wide one.
    """
    set_widths(ws, headers)
    if freeze:
        ws.freeze_panes = freeze
    if autofilter and headers:
        last = get_column_letter(len(headers))
        ws.auto_filter.ref = f"A1:{last}{max(row_count + 1, 1)}"


def flag_bad_status(
    ws: Worksheet, col: int, first_row: int, last_row: int, *, also_skip: Iterable[str] = (),
) -> None:
    """Red fill on any Status cell that is neither empty nor "ok".

    A conditional rule rather than a fill written per cell, so a reader who
    edits a Status by hand sees the colour follow, and the rows stay plain
    values that `parse_workbook` reads back unchanged.
    """
    if last_row < first_row:
        return
    letter = get_column_letter(col)
    ref = f"${letter}{first_row}"
    skip = "".join(f',{ref}<>"{s}"' for s in also_skip)
    ws.conditional_formatting.add(
        f"{letter}{first_row}:{letter}{last_row}",
        FormulaRule(
            formula=[f'AND(LEN({ref})>0,{ref}<>"ok"{skip})'],
            fill=BAD_FILL, font=BAD_FONT,
        ),
    )


def current_scale(ws: Worksheet, col: int, first_row: int, last_row: int) -> None:
    """Green-to-red over a current column: least current is the good end."""
    if last_row < first_row:
        return
    letter = get_column_letter(col)
    ws.conditional_formatting.add(
        f"{letter}{first_row}:{letter}{last_row}",
        ColorScaleRule(
            start_type="min", start_color="63BE7B",
            mid_type="percentile", mid_value=50, mid_color="FFEB84",
            end_type="max", end_color="F8696B",
        ),
    )


# --- Run sheet -------------------------------------------------------------

#: A label with this value is a section title, not a setting.
SECTION = None
#: Matched by name as well as by the empty value: a copied Run sheet (see
#: `rebuild_workbook`) can hold a genuinely empty setting, which must not turn
#: into a section heading.
SECTION_TITLES = {"Summary", "Rig", "Export"}


def write_run_entries(ws: Worksheet, entries: Sequence[tuple[Any, Any]]) -> None:
    """Label/value pairs, read top to bottom. A `(label, None)` whose label is a
    known section title is drawn as one; otherwise None is just an empty value."""
    write_header(ws, ["Setting", "Value"])
    for i, (label, value) in enumerate(entries, start=2):
        a = ws.cell(row=i, column=1, value=label)
        a.font = HEADER_FONT
        if value is SECTION and label in SECTION_TITLES:
            a.fill = SECTION_FILL
            ws.cell(row=i, column=2).fill = SECTION_FILL
            continue
        write_cell(ws, i, 2, value, str(label) if label is not None else None)
    set_widths(ws, ["Setting", "Value"])
    ws.freeze_panes = "A2"


def stamp(epoch: float | None) -> str:
    return "—" if epoch is None else time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(epoch))


def duration(started: float | None, finished: float | None) -> str:
    if started is None or finished is None or finished < started:
        return "—"
    s = int(round(finished - started))
    return f"{s // 3600}:{s % 3600 // 60:02d}:{s % 60:02d}"


def summary_entries(
    ok: Sequence[bool],
    measured: Iterable[float | None],
    measured_label: str,
    started: float | None = None,
    finished: float | None = None,
) -> list[tuple[str, Any]]:
    """The few numbers a reader wants before opening any data sheet."""
    readings = [m for m in measured if m is not None]
    return [
        ("Summary", SECTION),
        ("Points ok", sum(1 for x in ok if x)),
        ("Points failed", sum(1 for x in ok if not x)),
        (f"{measured_label} min [dBm]", min(readings) if readings else "—"),
        (f"{measured_label} max [dBm]", max(readings) if readings else "—"),
        ("Started", stamp(started)),
        ("Finished", stamp(finished)),
        ("Duration", duration(started, finished)),
    ]


def rig_entries(
    *,
    dut_mac: str | None,
    dut_name: str | None,
    power_sensor_idn: str | None,
    power_sensor_address: str | None,
    dc_analyzer_idn: str | None,
    dc_analyzer_address: str | None,
    dc_channel: int | None,
) -> list[tuple[str, Any]]:
    """Which bench produced the numbers. Without it two files from different
    sensors -- or a sensor swapped mid-week -- are indistinguishable."""
    def text(v: Any) -> Any:
        return "—" if v is None or v == "" else v

    return [
        ("Rig", SECTION),
        ("DUT MAC", text(dut_mac)),
        ("DUT device", text(dut_name)),
        ("Power sensor", text(power_sensor_idn)),
        ("Power sensor address", text(power_sensor_address)),
        ("DC analyzer", text(dc_analyzer_idn)),
        ("DC analyzer address", text(dc_analyzer_address)),
        ("DC channel", text(dc_channel)),
    ]


def export_entries(operator: str | None, notes: str | None) -> list[tuple[str, Any]]:
    entries: list[tuple[str, Any]] = [
        ("Export", SECTION),
        ("Exported", time.strftime("%Y-%m-%d %H:%M:%S")),
        ("App build", app_build()),
    ]
    # Only when given: a row reading "Operator —" on every file says nothing.
    if operator:
        entries.append(("Operator", operator))
    if notes:
        entries.append(("Notes", notes))
    return entries


@lru_cache(maxsize=1)
def app_build() -> str:
    """The git short SHA the backend is running from, or "unknown".

    Best-effort by design: the packaged build has no `.git` and often no git on
    PATH, and a missing version must never fail an export. Cached because it
    cannot change while the process runs -- and backend changes need a restart.
    """
    root = Path(__file__).resolve().parent.parent
    try:
        out = subprocess.run(
            ["git", "rev-parse", "--short", "HEAD"],
            cwd=root, capture_output=True, text=True, timeout=3,
            # No console window flashing up from the packaged (windowed) app.
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0) if sys.platform == "win32" else 0,
        )
    except (OSError, subprocess.SubprocessError):
        return "unknown"
    sha = out.stdout.strip()
    return sha if out.returncode == 0 and sha else "unknown"


# --- Filenames -------------------------------------------------------------

def mac_slug(mac: str | None) -> str:
    """Mirrors `macSlug` in frontend/src/lib/download.ts."""
    return (mac or "").replace(":", "").upper()


def export_filename(mac: str | None, test_slug: str, ext: str) -> str:
    """`<MAC or noDUT>_<test-slug>_<YYYY-MM-DDTHH-MM-SS>.<ext>`.

    The one scheme for every export, CSV and workbook alike; `exportName` in
    frontend/src/lib/download.ts builds the same name for the files the browser
    writes itself. Local time, as the Run sheet's own stamps are.
    """
    return f"{mac_slug(mac) or 'noDUT'}_{test_slug}_{time.strftime('%Y-%m-%dT%H-%M-%S')}.{ext}"
