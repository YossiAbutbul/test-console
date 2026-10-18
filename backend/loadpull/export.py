"""Excel export for Load Pull results.

Workbook layout:
- "All": every point, in the order it was measured.
- One sheet per frequency, holding a separate table per commanded power.

The split matters because a load pull contour only means anything at one
frequency and one drive level: stacking every combination in a single grid
leaves the reader to filter it back apart before the numbers can be compared.

Presentation (header style, two-decimal display, the red Status rule, the Run
sheet's sections) comes from `backend/workbook.py`, shared with the sweep.
"""
from __future__ import annotations

from io import BytesIO
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from .. import workbook as wbk
from .models import LoadPullMeta, LoadPullRow

#: Columns of the "All" sheet. The per-frequency sheets drop Freq and Set,
#: which are constant within each of their tables.
#:
#: "Measured [dBm]" and "X [Ω]" were "Power [dBm]" and "J [Ω]": the first so the
#: sweep, load-pull and CSV exports share one vocabulary, the second because
#: the column holds the reactance X, and "J" named the imaginary unit instead.
HEADERS = [
    "#", "Pos [mm]", "Pos [pulses]", "Freq [MHz]", "Set [dBm]", "Att [dB]",
    "Measured [dBm]", "Raw [dBm]", "CC [mA]", "R [Ω]", "X [Ω]", "S11 [dB]",
    "VSWR", "Status",
]
#: Earlier spellings, accepted on import so files written before a rename
#: still read back.
HEADER_ALIASES = {"Measured [dBm]": ("Power [dBm]",), "X [Ω]": ("J [Ω]",)}
#: Columns added after files were already in circulation. Matching by name
#: makes a missing one survivable, and an older workbook is still a complete
#: record of the run it holds, so their absence is not an error.
OPTIONAL_HEADERS = {"Raw [dBm]", "VSWR"}
#: Of those, the ones worked out from the columns beside them rather than
#: measured. Written but never read back, so a reader who edited the cell does
#: not get to contradict the measurement it came from.
DERIVED_HEADERS = {"VSWR"}
#: The per-frequency tables keep Att: unlike frequency and power it varies
#: *inside* one table, and it is what those rows are sorted by.
TABLE_HEADERS = [h for h in HEADERS if h not in ("Freq [MHz]", "Set [dBm]")]

TITLE_FONT = Font(bold=True, size=12)


def _milliamps(current_a: float | None) -> float | None:
    # Not rounded: the cell format shows two decimals, the cell keeps the
    # reading. See backend/workbook.py.
    return None if current_a is None else current_a * 1000.0


def _status(r: LoadPullRow) -> str:
    return r.error or "ok"


def _vswr(s11_db: float | None) -> float | None:
    """VSWR = (1 + |Γ|) / (1 - |Γ|), with |Γ| = 10^(S11/20).

    Derived here rather than carried on the row so the sheet cannot disagree
    with the S11 column printed next to it. The browser works the same figure
    out for its results table — see `vswr` in
    frontend/src/tests/loadPull/smith.ts, which is the definition this one
    tracks.

    |Γ| >= 1 is a total reflection, or in practice a stale calibration showing
    marginally more coming back than went out. The ratio is unbounded there, so
    the cell is left empty; the arithmetic left alone would turn the sign over
    and print a small number that reads like a good match.
    """
    if s11_db is None:
        return None
    gamma = 10.0 ** (s11_db / 20.0)
    if gamma >= 1.0:
        return None
    return (1.0 + gamma) / (1.0 - gamma)


def _values(idx: int, r: LoadPullRow, *, with_point: bool) -> list[Any]:
    """`with_point` includes Freq and Set — the two the sheet title already
    states on a per-frequency sheet."""
    head = [idx, r.pos_mm, r.pos_pulses]
    point = [r.freq_mhz, r.power_dbm_setting] if with_point else []
    # Attenuation stays on the per-frequency tables: unlike frequency and power
    # it varies *within* one of them, and it is what those rows are sorted by.
    tail = [
        r.att_db,
        # Raw sits beside the corrected figure: the correction moves the number
        # by tens of dB without moving where the sensor sat, and it is that
        # position in the meter's range which says how far to trust the point.
        r.power_dbm, r.power_dbm_raw,
        _milliamps(r.current_a), r.r_ohm, r.x_ohm, r.s11_db,
        _vswr(r.s11_db), _status(r),
    ]
    return head + point + tail


def _write_all(ws: Worksheet, rows: list[LoadPullRow]) -> None:
    wbk.write_header(ws, HEADERS)
    for i, r in enumerate(rows, start=1):
        for c, v in enumerate(_values(i, r, with_point=True), start=1):
            wbk.write_cell(ws, i + 1, c, v, HEADERS[c - 1])
    wbk.finish_table(ws, HEADERS, len(rows))
    wbk.flag_bad_status(ws, HEADERS.index("Status") + 1, 2, len(rows) + 1)


def _fmt_freq(mhz: float | None) -> str:
    if mhz is None:
        return "Unspecified"
    # 902.3 stays 902.3; 915.0 becomes 915, so the tab reads like the field did.
    return f"{mhz:g} MHz"


def _sheet_title(used: set[str], base: str) -> str:
    """Excel caps a tab at 31 characters and rejects duplicates."""
    title = base[:31]
    n = 2
    while title in used:
        suffix = f" ({n})"
        title = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(title)
    return title


def _write_frequency_sheet(ws: Worksheet, rows: list[LoadPullRow]) -> None:
    """One table per commanded power, stacked down the sheet."""
    powers = sorted({r.power_dbm_setting for r in rows}, key=lambda p: (p is None, p))
    row_at = 1
    for power in powers:
        # Sorted by attenuation, 0 first — the sweep records a whole trombone
        # cycle per setting, and a reader compares the cycles against each
        # other. Rows without one keep their measured order, ahead of any that
        # have it. Python's sort is stable, so position order survives within a
        # setting.
        group = sorted(
            (r for r in rows if r.power_dbm_setting == power),
            key=lambda r: (r.att_db is not None, r.att_db or 0.0),
        )
        label = "Unspecified" if power is None else f"{power:g} dBm"
        ws.cell(row=row_at, column=1, value=f"Set power: {label}").font = TITLE_FONT
        row_at += 1
        wbk.write_header(ws, TABLE_HEADERS, row=row_at)
        row_at += 1
        for i, r in enumerate(group, start=1):
            for c, v in enumerate(_values(i, r, with_point=False), start=1):
                wbk.write_cell(ws, row_at, c, v, TABLE_HEADERS[c - 1])
            row_at += 1
        # One blank line between tables, so a reader can tell where one ends.
        row_at += 1
    # No frozen header and no auto-filter: the tables are stacked, so there is
    # no single header row to pin, and Excel allows one filter per sheet. The
    # Status rule spans the whole column and skips the repeated header text.
    wbk.set_widths(ws, TABLE_HEADERS)
    wbk.flag_bad_status(
        ws, TABLE_HEADERS.index("Status") + 1, 1, max(row_at - 1, 1), also_skip=("Status",),
    )


def _float_or_none(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def parse_workbook(data: bytes) -> list[LoadPullRow]:
    """Read back a workbook this module wrote.

    Only the "All" sheet is read. The frequency tabs hold the same points split
    up, so importing those too would duplicate every row.
    """
    from openpyxl import load_workbook  # local: only importing pays for it

    wb = load_workbook(BytesIO(data), data_only=True)
    if "All" not in wb.sheetnames:
        raise ValueError(
            "no 'All' sheet - this is not a Load Pull workbook exported by this app"
        )
    ws = wb["All"]

    header = [c.value for c in ws[1]]
    # Matched by name, not by position: a workbook written before a column was
    # added is still a complete record of its own run, and refusing it would
    # strand every file exported up to that point. Derived columns are not
    # required at all -- they are recomputed from the measurements beside them.
    col = {name: i for i, name in enumerate(header) if isinstance(name, str)}
    for name, old in HEADER_ALIASES.items():
        for alias in old:
            if name not in col and alias in col:
                col[name] = col[alias]
    missing = [h for h in HEADERS if h not in col and h not in OPTIONAL_HEADERS]
    if missing:
        raise ValueError(
            f"unexpected columns: missing {missing}, found {header}. "
            "The file was exported by a different version of this app."
        )

    def cell(values: tuple[Any, ...], name: str) -> Any:
        i = col.get(name)
        if i is None:
            return None            # a column this file predates
        # A row saved from Excel can be shorter than the header if its trailing
        # cells were cleared.
        return values[i] if i < len(values) else None

    rows: list[LoadPullRow] = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        if cell(values, "#") is None:
            continue  # trailing blank row
        pos_mm = _float_or_none(cell(values, "Pos [mm]"))
        pos_pulses = _float_or_none(cell(values, "Pos [pulses]"))
        if pos_mm is None and pos_pulses is None:
            continue
        cc_ma = _float_or_none(cell(values, "CC [mA]"))
        status = str(cell(values, "Status") or "").strip()
        rows.append(LoadPullRow(
            pos_pulses=int(pos_pulses or 0),
            pos_mm=pos_mm or 0.0,
            freq_mhz=_float_or_none(cell(values, "Freq [MHz]")),
            power_dbm_setting=_float_or_none(cell(values, "Set [dBm]")),
            att_db=_float_or_none(cell(values, "Att [dB]")),
            power_dbm=_float_or_none(cell(values, "Measured [dBm]")),
            power_dbm_raw=_float_or_none(cell(values, "Raw [dBm]")),
            current_a=None if cc_ma is None else cc_ma / 1000.0,
            r_ohm=_float_or_none(cell(values, "R [Ω]")),
            x_ohm=_float_or_none(cell(values, "X [Ω]")),
            s11_db=_float_or_none(cell(values, "S11 [dB]")),
            # Anything that is not a plain "ok" was the reason the row has no
            # numbers, so it travels back as the error.
            error=None if status.lower() == "ok" else (status or None),
        ))
    return rows


#: Spelt as the CSV exports and the app spell it -- see the sweep's copy.
_PA_MODE_NAMES = {0: "Off", 1: "On", 2: "Auto"}


def _mm(pulses: int | None) -> str:
    return "—" if pulses is None else f"{pulses / 400:.2f} mm ({pulses} pulses)"


def _write_run_sheet(ws: Worksheet, meta: LoadPullMeta, rows: list[LoadPullRow]) -> None:
    """Label/value pairs — read top to bottom, not a table."""
    def dash(v: Any) -> Any:
        return "—" if v is None else v

    entries: list[tuple[str, Any]] = [
        ("Points", len(rows)),
        ("Frequencies", meta.freq_spec or "—"),
        ("Powers", meta.power_spec or "—"),
        ("PA mode", f"{_PA_MODE_NAMES.get(meta.pa_mode, '?')} ({meta.pa_mode})"
                    if meta.pa_mode is not None else "—"),
        ("Settle [ms]", dash(meta.settle_ms)),
        ("Trombone zero", _mm(meta.zero_pulses)),
        ("Trombone end", _mm(meta.end_pulses)),
        ("Delta X [mm]", dash(meta.delta_x_mm)),
        ("Path loss default [dB]", dash(meta.path_loss_default_db)),
    ]
    # Per frequency, so a reader can see which figure corrected which column -
    # and which of them was a guess.
    for pt in meta.path_loss_points:
        label = f"Path loss @ {pt.freq_mhz:g} MHz [dB]"
        entries.append((label, f"{pt.db:g}" + ("" if pt.calibrated else "  (uncalibrated)")))
    entries += wbk.summary_entries(
        [r.error is None for r in rows], (r.power_dbm for r in rows), "Measured",
        meta.started_at, meta.finished_at,
    )
    entries += wbk.rig_entries(
        dut_mac=meta.dut_mac, dut_name=meta.dut_name,
        power_sensor_idn=meta.power_sensor_idn, power_sensor_address=meta.power_sensor_address,
        dc_analyzer_idn=meta.dc_analyzer_idn, dc_analyzer_address=meta.dc_analyzer_address,
        dc_channel=meta.dc_channel,
    )
    entries += wbk.export_entries(meta.operator, meta.notes)
    wbk.write_run_entries(ws, entries)


def build_workbook(
    rows: Iterable[LoadPullRow], meta: LoadPullMeta | None = None,
) -> bytes:
    """Render Load Pull rows to an .xlsx file, returned as bytes."""
    rows = list(rows)
    wb = Workbook()

    ws_all = wb.active
    ws_all.title = "All"
    _write_all(ws_all, rows)

    # Right after "All" rather than last: appended behind a dozen frequency
    # tabs it would not be found.
    if meta is not None:
        _write_run_sheet(wb.create_sheet(title="Run", index=1), meta, rows)

    # Ascending, with unspecified last so a real frequency is never buried.
    freqs = sorted({r.freq_mhz for r in rows}, key=lambda f: (f is None, f))
    used = {"All", "Run"}
    for f in freqs:
        subset = [r for r in rows if r.freq_mhz == f]
        _write_frequency_sheet(wb.create_sheet(title=_sheet_title(used, _fmt_freq(f))), subset)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
